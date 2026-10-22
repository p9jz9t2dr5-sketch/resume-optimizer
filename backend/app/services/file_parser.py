"""文件解析：按扩展名校验并提取正文。

- PDF → pypdf；没有文字层时渲染各页并用 Qwen-VL OCR
  （扫描件的页面图会先发给视觉模型，PII 脱敏发生在 OCR 之后）
- DOCX → python-docx
- TXT → 直接读取
- 图片 → 交给 llm_service 的 Qwen-VL OCR
"""

import asyncio
import logging
import os
import tempfile
from typing import Optional
from pypdf import PdfReader
from docx import Document

from app.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()

# 扫描件通常只有页码/元数据这几个字符。用「总字数下限 + 每页下限」双阈值，
# 避免把一份「页数多但文字少」的扫描件误判成纯文字 PDF。
PDF_TEXT_LAYER_MIN_CHARS = 10
PDF_TEXT_LAYER_MIN_CHARS_PER_PAGE = 10
# Protect the model budget from accidentally uploading a book/portfolio.
PDF_OCR_MAX_PAGES = 8
# Keep page images large enough for Chinese OCR without sending enormous images.
PDF_OCR_TARGET_LONG_EDGE = 1800


def parse_pdf(file_path: str) -> str:
    """Extract text from PDF file."""
    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _meaningful_char_count(text: str) -> int:
    """Count visible characters so whitespace-only PDFs are treated as scanned."""
    return sum(1 for char in text if not char.isspace())


def _pdf_page_count(file_path: str) -> int:
    """页数只用于决定是否需要 OCR；读不出来就按 1 页处理。"""
    try:
        return max(1, len(PdfReader(file_path).pages))
    except Exception:
        return 1


def _render_pdf_pages(file_path: str, temp_dir: str, max_pages: int) -> tuple[list[str], int]:
    """同步渲染 PDF 页面（由 asyncio.to_thread 调用）。

    ``pypdfium2`` 的渲染是 CPU 密集的同步调用；直接在协程里跑会阻塞整个事件循环，
    2 核机器上 8 页就可能把同一 worker 的其他请求一起卡住。
    """
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(file_path)
    image_paths: list[str] = []
    try:
        page_count = len(pdf)
        if page_count == 0:
            raise RuntimeError("PDF 中没有页面")

        for page_index in range(min(page_count, max_pages)):
            page = pdf[page_index]
            try:
                width, height = page.get_size()
                long_edge = max(float(width), float(height), 1.0)
                scale = min(2.5, PDF_OCR_TARGET_LONG_EDGE / long_edge)
                scale = max(scale, 0.5)

                bitmap = page.render(scale=scale)
                try:
                    image = bitmap.to_pil().convert("RGB")
                finally:
                    bitmap.close()

                image_path = os.path.join(temp_dir, f"page-{page_index + 1}.jpg")
                image.save(image_path, format="JPEG", quality=90, optimize=True)
                image_paths.append(image_path)
            finally:
                page.close()
    finally:
        pdf.close()
    return image_paths, page_count


async def _ocr_pdf(file_path: str, fallback_text: str = "") -> str:
    """Render a text-less PDF page by page and transcribe each page with Qwen-VL.

    ``pypdf`` only reads a PDF's text layer. Image-only/scanned PDFs therefore
    return an empty string even though the pages are visible. Rendering each
    page to JPEG lets us reuse the existing vision OCR path.
    """
    if not settings.DASHSCOPE_API_KEY:
        if fallback_text.strip():
            logger.warning("scanned PDF has no text layer and DASHSCOPE_API_KEY is unset")
            return fallback_text
        raise RuntimeError(
            "PDF 没有可提取的文字层，且未配置 DASHSCOPE_API_KEY，无法进行视觉 OCR"
        )

    from app.services.llm_service import llm_service

    page_texts: list[str] = []
    with tempfile.TemporaryDirectory(prefix="resume-pdf-ocr-") as temp_dir:
        try:
            image_paths, page_count = await asyncio.to_thread(
                _render_pdf_pages, file_path, temp_dir, PDF_OCR_MAX_PAGES
            )
        except Exception as exc:
            raise RuntimeError(f"PDF 页面渲染失败：{type(exc).__name__}: {exc}") from exc

        if page_count > PDF_OCR_MAX_PAGES:
            logger.warning(
                "scanned PDF has %d pages; only the first %d will be OCR'd",
                page_count,
                PDF_OCR_MAX_PAGES,
            )

        for page_index, image_path in enumerate(image_paths):
            try:
                page_text = await llm_service.parse_image(image_path)
            except Exception as exc:
                logger.warning(
                    "PDF OCR failed for page %d: %s: %s",
                    page_index + 1,
                    type(exc).__name__,
                    exc,
                )
                continue

            if page_text.strip():
                page_texts.append(f"## 第 {page_index + 1} 页\n\n{page_text.strip()}")

    if page_texts:
        return "\n\n".join(page_texts)
    if fallback_text.strip():
        logger.warning("PDF OCR returned no text; keeping the small extracted text layer")
        return fallback_text
    raise RuntimeError("PDF 没有可提取的文字层，视觉 OCR 也未识别到内容")


def parse_docx(file_path: str) -> str:
    """Extract text from Word document."""
    doc = Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    return "\n".join(text_parts)


async def parse_file(file_path: str) -> str:
    """Parse file based on extension. Returns extracted text."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        text = parse_pdf(file_path)
        chars = _meaningful_char_count(text)
        required = max(
            PDF_TEXT_LAYER_MIN_CHARS,
            _pdf_page_count(file_path) * PDF_TEXT_LAYER_MIN_CHARS_PER_PAGE,
        )
        if chars >= required:
            return text
        logger.info(
            "PDF has little/no text layer (%d visible chars < %d); falling back to page OCR",
            chars,
            required,
        )
        return await _ocr_pdf(file_path, fallback_text=text)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    elif ext in (".png", ".jpg", ".jpeg", ".webp"):
        # 图片简历：调用视觉模型做 OCR（延迟导入，避免循环依赖）
        from app.services.llm_service import llm_service
        return await llm_service.parse_image(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".png", ".jpg", ".jpeg", ".webp"}


def allowed_file(filename: str) -> bool:
    """扩展名是否在白名单里（PDF / DOCX / DOC / TXT / PNG / JPG / JPEG / WEBP）。"""
    return os.path.splitext(filename)[1].lower() in ALLOWED_EXTENSIONS
